import os
import openpyxl
from openpyxl.styles import Font, PatternFill

def export_to_excel(successful_records, pending_records, output_path):
    """
    Exports the records to an Excel workbook with two sheets: Despesas and Pendencias.
    """
    wb = openpyxl.Workbook()
    
    # 1. Main Sheet: Despesas
    ws_despesas = wb.active
    ws_despesas.title = "Despesas"
    
    headers_despesas = ["DATA DO PAGAMENTO", "DESCRIÇÃO DA DESPESA", "VALOR DA DESPESA", "STATUS"]
    ws_despesas.append(headers_despesas)
    
    # Sort successful records by month, date, description
    sorted_success = sorted(
        successful_records, 
        key=lambda x: (x.get('month', ''), str(x.get('date', '')), x.get('description', ''))
    )
    
    yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    for row_idx, rec in enumerate(sorted_success, start=2):
        row_data = [rec['date'], rec['description'], rec['value'], rec['status']]
        ws_despesas.append(row_data)
        
        # Apply yellow fill if status is not OK
        if rec['status'] != "OK":
            for col_idx in range(1, 5):
                ws_despesas.cell(row=row_idx, column=col_idx).fill = yellow_fill
                
    # Format the columns
    for row_idx in range(2, len(sorted_success) + 2):
        # Format Date
        cell_date = ws_despesas.cell(row=row_idx, column=1)
        if cell_date.value:
            cell_date.number_format = 'DD/MM/YYYY'
            
        # Format Value
        cell_val = ws_despesas.cell(row=row_idx, column=3)
        if cell_val.value != "":
            cell_val.number_format = '"R$" #,##0.00'

    # Make headers bold
    for cell in ws_despesas[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        
    # Adjust column widths
    ws_despesas.column_dimensions['A'].width = 25
    ws_despesas.column_dimensions['B'].width = 60
    ws_despesas.column_dimensions['C'].width = 25
    ws_despesas.column_dimensions['D'].width = 40

    # 2. Secondary Sheet: Pendencias
    ws_pendencias = wb.create_sheet(title="Pendencias")
    headers_pendencias = ["MÊS", "ARQUIVO_ORIGEM", "CAMINHO_COMPLETO", "STATUS", "MOTIVO"]
    ws_pendencias.append(headers_pendencias)
    
    # Sort pending records by month, filename
    sorted_pending = sorted(
        pending_records,
        key=lambda x: (x.get('month', ''), x.get('filename', ''))
    )
    
    for rec in sorted_pending:
        row = [
            rec.get('month', ''),
            rec.get('filename', ''),
            rec.get('filepath', ''),
            rec.get('status', 'Pendente'),
            rec.get('motivo', '')
        ]
        ws_pendencias.append(row)
        
    for cell in ws_pendencias[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
        
    ws_pendencias.column_dimensions['A'].width = 15
    ws_pendencias.column_dimensions['B'].width = 40
    ws_pendencias.column_dimensions['C'].width = 80
    ws_pendencias.column_dimensions['D'].width = 15
    ws_pendencias.column_dimensions['E'].width = 50

    wb.save(output_path)
